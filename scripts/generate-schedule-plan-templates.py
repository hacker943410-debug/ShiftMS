from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidationList


ROOT = Path(__file__).resolve().parent.parent
SAMPLE_DIR = ROOT / "양식샘플"


def clear_range(ws, start_row, end_row, start_col, end_col):
    for row in ws.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=start_col,
        max_col=end_col,
    ):
        for cell in row:
            if type(cell).__name__ == "MergedCell":
                continue
            cell.value = None


def keep_data_validation(formula):
    if not formula:
        return True

    blocked_tokens = [
        "김규민",
        "BP(",
        "정현우",
        "김영서",
        "이동희",
        "고성은",
        "박재환",
        "정혜진",
        "박준하",
        "이상우",
        "최종수",
        "이진영",
        "임대현",
        "전인환",
        "문승건",
        "김무성",
        "유호진",
        "김남영",
        "박정수",
        "조예덕",
        "최의헌",
        "이준범",
        "배상현",
        "강태경",
        "문경철",
        "김성수",
        "풍기웅",
        "김희경",
        "김학성",
    ]

    return not any(token in formula for token in blocked_tokens)


def rebuild_validations(worksheet):
    kept_validations = []
    for validation in worksheet.data_validations.dataValidation:
        if keep_data_validation(validation.formula1):
            kept_validations.append(validation)

    worksheet.data_validations = DataValidationList()
    for validation in kept_validations:
        worksheet.data_validations.append(validation)


def build_template(config):
    workbook = load_workbook(config["sample_path"])
    worksheet = workbook[workbook.sheetnames[0]]

    for address in config["clear_cells"]:
        worksheet[address] = None

    for start_row, end_row, start_col, end_col in config["clear_ranges"]:
        clear_range(worksheet, start_row, end_row, start_col, end_col)

    rebuild_validations(worksheet)
    workbook.save(config["output_path"])


def main():
    build_template(
        {
            "sample_path": SAMPLE_DIR / "근무표_샘플1.xlsx",
            "output_path": SAMPLE_DIR / "근무표_템플릿1.xlsx",
            "clear_cells": ["C3", "B7", "W6", "AY8", "BK8", "BK30"],
            "clear_ranges": [
                (9, 44, 3, 23),
                (12, 42, 25, 51),
                (11, 26, 53, 63),
                (34, 40, 53, 63),
            ],
        }
    )
    build_template(
        {
            "sample_path": SAMPLE_DIR / "근무표_샘플2.xlsx",
            "output_path": SAMPLE_DIR / "근무표_템플릿2.xlsx",
            "clear_cells": ["C3", "B7", "W6", "BG8", "BS8", "BS30"],
            "clear_ranges": [
                (9, 68, 3, 23),
                (12, 42, 25, 59),
                (11, 26, 61, 71),
                (34, 40, 61, 71),
            ],
        }
    )


if __name__ == "__main__":
    main()
